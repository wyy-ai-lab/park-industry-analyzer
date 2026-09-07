"""
园区产业诊断报告导出
支持 Markdown、Word、PDF、HTML 四种格式。
"""

import base64
import io
from datetime import datetime
from typing import Any, Dict, List, Optional

from engine.report_export import _find_cjk_font


STATUS_COLOR_HEX = {
    "强势": "#0056b3",
    "正常": "#34c759",
    "薄弱": "#ff9500",
    "缺失": "#ff3b30",
}


def _today() -> str:
    return datetime.now().strftime("%Y年%m月%d日")


def _today_compact() -> str:
    return datetime.now().strftime("%Y%m%d")


def _safe_get(obj: Dict[str, Any], key: str, default: Any = "") -> Any:
    return obj.get(key, default) if isinstance(obj, dict) else default


def _fmt_list(items: List[str], empty_text: str = "无") -> str:
    return "、".join(items) if items else empty_text


def _fmt_number(n) -> str:
    try:
        return f"{n:,.1f}" if isinstance(n, float) else f"{n:,}"
    except Exception:
        return str(n)


def _segment_status_table(metrics: Dict[str, Any]) -> List[Dict[str, str]]:
    """生成产业链环节状态表格数据"""
    from engine.chain_position import SEGMENT_LAYER_MAP

    segment_dist = metrics.get("segment_distribution", {})
    segment_status = metrics.get("segment_strength", {})
    rows = []
    for seg, info in segment_dist.items():
        status = "正常"
        if seg in segment_status.get("strong", []):
            status = "强势"
        elif seg in segment_status.get("weak", []):
            status = "薄弱"
        elif seg in segment_status.get("missing", []):
            status = "缺失"
        rows.append({
            "环节": seg,
            "层级": info.get("chain_position", "—"),
            "企业数": str(info.get("count", 0)),
            "年产值": f"{info.get('revenue', 0):.2f} 亿元",
            "状态": status,
        })
    # 缺失环节也加入表格
    for seg in segment_status.get("missing", []):
        if seg not in segment_dist:
            layer, _ = SEGMENT_LAYER_MAP.get(seg, ("—", ""))
            rows.append({
                "环节": seg,
                "层级": layer,
                "企业数": "0",
                "年产值": "0.00 亿元",
                "状态": "缺失",
            })
    return rows


def build_park_markdown_report(metrics: Dict[str, Any], diagnosis: Dict[str, Any]) -> str:
    """生成园区产业诊断 Markdown 报告"""
    park_name = metrics.get("park_name", "未知园区")
    totals = metrics.get("totals", {})
    date_str = _today()

    md = f"""# {park_name} 产业诊断报告

<div align="center">

**报告生成日期**：{date_str}

</div>

---

## 一、核心指标

| 指标 | 数值 |
|------|------|
| 企业总数 | {totals.get('enterprise_count', 0)} 家 |
| 年产值 | {totals.get('total_revenue', 0):.2f} 亿元 |
| 员工总数 | {totals.get('total_employees', 0):,} 人 |
| 高新技术企业 | {totals.get('high_tech_count', 0)} 家 |
| 小巨人企业 | {totals.get('little_giant_count', 0)} 家 |
| 产业链完整度评分 | {metrics.get('completeness_score', 0)} 分 |
| 本地配套率 | {metrics.get('local_support_rate', 0)}% |

---

## 二、产业分布与产业链层级

### 产业领域分布

| 产业领域 | 企业数 |
|----------|--------|
"""
    for k, v in metrics.get("sub_industry_distribution", {}).items():
        md += f"| {k} | {v} 家 |\n"

    md += """
### 产业链层级分布

| 层级 | 企业数 |
|------|--------|
"""
    for k, v in metrics.get("chain_distribution", {}).items():
        md += f"| {k} | {v} 家 |\n"

    md += """
---

## 三、企业梯队金字塔

| 梯队 | 企业数 |
|------|--------|
"""
    tier_order = ["链主企业", "骨干企业", "高新技术企业", "科技型中小企业", "配套服务企业"]
    for k in tier_order:
        v = metrics.get("tier_distribution", {}).get(k, 0)
        if v > 0:
            md += f"| {k} | {v} 家 |\n"

    md += """
---

## 四、产业链强弱分析

### 强势环节
"""
    strong = metrics.get("segment_strength", {}).get("strong", [])
    md += (_fmt_list(strong) + "\n\n") if strong else "暂无\n\n"

    md += "### 薄弱环节\n"
    weak = metrics.get("segment_strength", {}).get("weak", [])
    md += (_fmt_list(weak) + "\n\n") if weak else "暂无\n\n"

    md += "### 缺失环节\n"
    missing = metrics.get("segment_strength", {}).get("missing", [])
    md += (_fmt_list(missing) + "\n\n") if missing else "暂无\n\n"

    md += "### 风险环节\n"
    risk = metrics.get("segment_strength", {}).get("risk", [])
    md += (_fmt_list(risk) + "\n\n") if risk else "暂无\n\n"

    md += """### 环节明细

| 环节 | 层级 | 企业数 | 年产值 | 状态 |
|------|------|--------|--------|------|
"""
    for row in _segment_status_table(metrics):
        md += f"| {row['环节']} | {row['层级']} | {row['企业数']} | {row['年产值']} | {row['状态']} |\n"

    md += f"""
---

## 五、本地配套率与产业链完整度

- **本地配套率**：{metrics.get('local_support_rate', 0)}%
- **产业链完整度评分**：{metrics.get('completeness_score', 0)} 分

---

## 六、发展建议

### 整体判断

{diagnosis.get('overall_assessment', '暂无')}

### 核心优势
"""
    for item in diagnosis.get("core_strengths", []):
        md += f"- {item}\n"

    md += "\n### 关键短板\n"
    for item in diagnosis.get("key_gaps", []):
        md += f"- {item}\n"

    md += "\n### 招商补链建议\n"
    for item in diagnosis.get("investment_suggestions", []):
        md += f"- {item}\n"

    md += "\n### 企业梯度培育建议\n"
    for item in diagnosis.get("cultivation_suggestions", []):
        md += f"- {item}\n"

    md += "\n### 政策建议\n"
    for item in diagnosis.get("policy_suggestions", []):
        md += f"- {item}\n"

    md += f"""
### 风险提醒

{diagnosis.get('risk_warning', '暂无')}

---

## 七、TOP10 头部企业

| 排名 | 企业名称 | 产业领域 | 年产值（亿元） |
|------|----------|----------|----------------|
"""
    for i, ent in enumerate(metrics.get("top_enterprises", [])[:10], 1):
        md += f"| {i} | {ent.get('name', '—')} | {ent.get('sub_industry', '—')} | {ent.get('annual_revenue', 0):.2f} |\n"

    md += """
---

*本报告由园区产业分析智能体自动生成*
"""
    return md


def build_park_word_report(
    metrics: Dict[str, Any],
    diagnosis: Dict[str, Any],
    charts_bytes: Dict[str, bytes],
) -> io.BytesIO:
    """生成园区产业诊断 Word 报告，返回 BytesIO"""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml

    park_name = metrics.get("park_name", "未知园区")
    totals = metrics.get("totals", {})
    date_str = _today()

    doc = Document()

    # 默认字体
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(10.5)

    # 封面
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title.add_run(f"{park_name}\n产业诊断报告")
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(29, 29, 31)

    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = info.add_run(f"报告生成日期：{date_str}")
    run.font.size = Pt(14)

    doc.add_page_break()

    # 一、核心指标
    doc.add_heading("一、核心指标", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "指标"
    hdr[1].text = "数值"
    for cell in hdr:
        cell.paragraphs[0].runs[0].font.bold = True

    rows = [
        ("企业总数", f"{totals.get('enterprise_count', 0)} 家"),
        ("年产值", f"{totals.get('total_revenue', 0):.2f} 亿元"),
        ("员工总数", f"{totals.get('total_employees', 0):,} 人"),
        ("高新技术企业", f"{totals.get('high_tech_count', 0)} 家"),
        ("小巨人企业", f"{totals.get('little_giant_count', 0)} 家"),
        ("产业链完整度评分", f"{metrics.get('completeness_score', 0)} 分"),
        ("本地配套率", f"{metrics.get('local_support_rate', 0)}%"),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    # 二、产业分布与产业链层级
    doc.add_heading("二、产业分布与产业链层级", level=1)
    doc.add_heading("产业领域分布", level=2)
    _add_chart_or_placeholder(doc, charts_bytes.get("industry_pie"))
    doc.add_heading("产业链层级分布", level=2)
    _add_chart_or_placeholder(doc, charts_bytes.get("chain_layer"))

    # 三、企业梯队金字塔
    doc.add_heading("三、企业梯队金字塔", level=1)
    _add_chart_or_placeholder(doc, charts_bytes.get("tier_pyramid"))

    # 四、产业链强弱分析
    doc.add_heading("四、产业链强弱分析", level=1)
    for title, key in [
        ("强势环节", "strong"),
        ("薄弱环节", "weak"),
        ("缺失环节", "missing"),
        ("风险环节", "risk"),
    ]:
        doc.add_heading(title, level=2)
        items = metrics.get("segment_strength", {}).get(key, [])
        if items:
            for item in items:
                doc.add_paragraph(item, style="List Bullet")
        else:
            doc.add_paragraph("暂无")

    doc.add_heading("环节明细", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    headers = ["环节", "层级", "企业数", "年产值", "状态"]
    for cell, text in zip(hdr, headers):
        cell.text = text
        cell.paragraphs[0].runs[0].font.bold = True

    for row in _segment_status_table(metrics):
        cells = table.add_row().cells
        cells[0].text = row["环节"]
        cells[1].text = row["层级"]
        cells[2].text = row["企业数"]
        cells[3].text = row["年产值"]
        cells[4].text = row["状态"]
        color = STATUS_COLOR_HEX.get(row["状态"])
        if color:
            _set_docx_cell_shading(cells[4], color)

    # 五、本地配套率与产业链完整度
    doc.add_heading("五、本地配套率与产业链完整度", level=1)
    doc.add_paragraph(f"本地配套率：{metrics.get('local_support_rate', 0)}%")
    doc.add_paragraph(f"产业链完整度评分：{metrics.get('completeness_score', 0)} 分")

    # 六、发展建议
    doc.add_heading("六、发展建议", level=1)
    doc.add_heading("整体判断", level=2)
    doc.add_paragraph(diagnosis.get("overall_assessment", "暂无"))

    for title, key in [
        ("核心优势", "core_strengths"),
        ("关键短板", "key_gaps"),
        ("招商补链建议", "investment_suggestions"),
        ("企业梯度培育建议", "cultivation_suggestions"),
        ("政策建议", "policy_suggestions"),
    ]:
        doc.add_heading(title, level=2)
        items = diagnosis.get(key, [])
        if items:
            for item in items:
                doc.add_paragraph(item, style="List Bullet")
        else:
            doc.add_paragraph("暂无")

    doc.add_heading("风险提醒", level=2)
    doc.add_paragraph(diagnosis.get("risk_warning", "暂无"))

    # 七、TOP10 头部企业
    doc.add_heading("七、TOP10 头部企业", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    headers = ["排名", "企业名称", "产业领域", "年产值（亿元）"]
    for cell, text in zip(hdr, headers):
        cell.text = text
        cell.paragraphs[0].runs[0].font.bold = True

    for i, ent in enumerate(metrics.get("top_enterprises", [])[:10], 1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = ent.get("name", "—")
        cells[2].text = ent.get("sub_industry", "—")
        cells[3].text = f"{ent.get('annual_revenue', 0):.2f}"

    # 页脚
    doc.add_paragraph()
    footer = doc.add_paragraph("本报告由园区产业分析智能体自动生成")
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(107, 114, 128)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_park_pdf_report(
    metrics: Dict[str, Any],
    diagnosis: Dict[str, Any],
    charts_bytes: Dict[str, bytes],
) -> io.BytesIO:
    """生成园区产业诊断 PDF 报告，返回 BytesIO"""
    from fpdf import FPDF

    font_path = _find_cjk_font()
    if not font_path:
        raise RuntimeError(
            "未找到中文字体，PDF 导出无法显示中文。"
            "请在 Windows 系统字体目录保留 simsun、simhei 或 msyh 字体之一。"
        )

    park_name = metrics.get("park_name", "未知园区")
    totals = metrics.get("totals", {})
    date_str = _today()

    pdf = FPDF()
    pdf.add_font("cn", "", font_path, uni=True)
    pdf.add_font("cn", "B", font_path, uni=True)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_left_margin(15)
    pdf.set_right_margin(15)

    # 封面
    pdf.add_page()
    pdf.set_font("cn", "B", 26)
    pdf.cell(0, 20, f"{park_name}", ln=True, align="C")
    pdf.cell(0, 14, "产业诊断报告", ln=True, align="C")
    pdf.ln(10)
    pdf.set_font("cn", "", 14)
    pdf.cell(0, 10, f"报告生成日期：{date_str}", ln=True, align="C")

    # 一、核心指标
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "一、核心指标", ln=True)
    pdf.ln(2)

    col_widths = [65, 60]
    line_height = 7
    pdf.set_font("cn", "B", 10)
    _pdf_table_row(pdf, ["指标", "数值"], col_widths, line_height, aligns=["C", "C"])
    pdf.set_font("cn", "", 10)
    rows = [
        ("企业总数", f"{totals.get('enterprise_count', 0)} 家"),
        ("年产值", f"{totals.get('total_revenue', 0):.2f} 亿元"),
        ("员工总数", f"{totals.get('total_employees', 0):,} 人"),
        ("高新技术企业", f"{totals.get('high_tech_count', 0)} 家"),
        ("小巨人企业", f"{totals.get('little_giant_count', 0)} 家"),
        ("产业链完整度评分", f"{metrics.get('completeness_score', 0)} 分"),
        ("本地配套率", f"{metrics.get('local_support_rate', 0)}%"),
    ]
    for label, value in rows:
        _pdf_table_row(pdf, [label, value], col_widths, line_height, aligns=["L", "C"])

    # 二、产业分布与产业链层级
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "二、产业分布与产业链层级", ln=True)
    pdf.ln(2)
    _pdf_insert_chart(pdf, charts_bytes.get("industry_pie"), title="产业领域分布")
    _pdf_insert_chart(pdf, charts_bytes.get("chain_layer"), title="产业链层级分布")

    # 三、企业梯队金字塔
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "三、企业梯队金字塔", ln=True)
    pdf.ln(2)
    _pdf_insert_chart(pdf, charts_bytes.get("tier_pyramid"))

    # 四、产业链强弱分析
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "四、产业链强弱分析", ln=True)
    pdf.ln(2)

    for title, key in [
        ("强势环节", "strong"),
        ("薄弱环节", "weak"),
        ("缺失环节", "missing"),
        ("风险环节", "risk"),
    ]:
        pdf.set_font("cn", "B", 12)
        pdf.cell(0, 8, title, ln=True)
        pdf.set_font("cn", "", 10)
        items = metrics.get("segment_strength", {}).get(key, [])
        if items:
            _pdf_safe_multi_cell(pdf, "、".join(items))
        else:
            _pdf_safe_multi_cell(pdf, "暂无")
        pdf.ln(2)

    pdf.set_font("cn", "B", 12)
    pdf.cell(0, 8, "环节明细", ln=True)
    col_widths = [35, 25, 25, 40, 25]
    pdf.set_font("cn", "B", 9)
    _pdf_table_row(
        pdf,
        ["环节", "层级", "企业数", "年产值", "状态"],
        col_widths,
        6,
        aligns=["C", "C", "C", "C", "C"],
    )
    pdf.set_font("cn", "", 9)
    for row in _segment_status_table(metrics):
        fill = None
        if row["状态"] in STATUS_COLOR_HEX:
            fill = _hex_to_rgb(STATUS_COLOR_HEX[row["状态"]])
        _pdf_table_row(
            pdf,
            [row["环节"], row["层级"], row["企业数"], row["年产值"], row["状态"]],
            col_widths,
            6,
            aligns=["L", "C", "C", "R", "C"],
            fills=[None, None, None, None, fill],
        )

    # 五、本地配套率与产业链完整度
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "五、本地配套率与产业链完整度", ln=True)
    pdf.ln(2)
    pdf.set_font("cn", "", 11)
    _pdf_safe_multi_cell(pdf, f"本地配套率：{metrics.get('local_support_rate', 0)}%")
    _pdf_safe_multi_cell(pdf, f"产业链完整度评分：{metrics.get('completeness_score', 0)} 分")

    # 六、发展建议
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "六、发展建议", ln=True)
    pdf.ln(2)

    pdf.set_font("cn", "B", 12)
    pdf.cell(0, 8, "整体判断", ln=True)
    pdf.set_font("cn", "", 10)
    _pdf_safe_multi_cell(pdf, diagnosis.get("overall_assessment", "暂无"))
    pdf.ln(2)

    for title, key in [
        ("核心优势", "core_strengths"),
        ("关键短板", "key_gaps"),
        ("招商补链建议", "investment_suggestions"),
        ("企业梯度培育建议", "cultivation_suggestions"),
        ("政策建议", "policy_suggestions"),
    ]:
        pdf.set_font("cn", "B", 12)
        pdf.cell(0, 8, title, ln=True)
        pdf.set_font("cn", "", 10)
        items = diagnosis.get(key, [])
        if items:
            for item in items:
                _pdf_safe_multi_cell(pdf, "- " + item)
        else:
            _pdf_safe_multi_cell(pdf, "暂无")
        pdf.ln(2)

    pdf.set_font("cn", "B", 12)
    pdf.cell(0, 8, "风险提醒", ln=True)
    pdf.set_font("cn", "", 10)
    _pdf_safe_multi_cell(pdf, diagnosis.get("risk_warning", "暂无"))

    # 七、TOP10 头部企业
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "七、TOP10 头部企业", ln=True)
    pdf.ln(2)

    col_widths = [18, 70, 45, 45]
    pdf.set_font("cn", "B", 9)
    _pdf_table_row(
        pdf,
        ["排名", "企业名称", "产业领域", "年产值（亿元）"],
        col_widths,
        6.5,
        aligns=["C", "C", "C", "C"],
    )
    pdf.set_font("cn", "", 9)
    for i, ent in enumerate(metrics.get("top_enterprises", [])[:10], 1):
        _pdf_table_row(
            pdf,
            [str(i), ent.get("name", "—"), ent.get("sub_industry", "—"), f"{ent.get('annual_revenue', 0):.2f}"],
            col_widths,
            6.5,
            aligns=["C", "L", "C", "R"],
        )

    # 页脚
    pdf.ln(8)
    pdf.set_font("cn", "", 9)
    pdf.set_text_color(128, 128, 128)
    pdf.cell(0, 8, "本报告由园区产业分析智能体自动生成", ln=True, align="C")

    buffer = io.BytesIO(bytes(pdf.output(dest="S")))
    return buffer


def build_park_html_report(
    metrics: Dict[str, Any],
    diagnosis: Dict[str, Any],
    charts_base64: Dict[str, str],
) -> str:
    """生成园区产业诊断 HTML 报告，图表使用 base64 内嵌"""
    park_name = metrics.get("park_name", "未知园区")
    totals = metrics.get("totals", {})
    date_str = _today()

    def _img(key: str, alt: str) -> str:
        b64 = charts_base64.get(key)
        if not b64:
            return f""
        return f'<img src="data:image/png;base64,{b64}" alt="{alt}" style="max-width:100%;height:auto;border-radius:12px;" />'

    # 核心指标卡片
    metric_cards = "\n".join(
        f"""
      <div class="metric-card">
        <div class="metric-value">{value}</div>
        <div class="metric-label">{label}</div>
      </div>"""
        for label, value in [
            ("企业总数", f"{totals.get('enterprise_count', 0)} 家"),
            ("年产值", f"{totals.get('total_revenue', 0):.2f} 亿元"),
            ("员工总数", f"{totals.get('total_employees', 0):,} 人"),
            ("高新技术企业", f"{totals.get('high_tech_count', 0)} 家"),
            ("小巨人企业", f"{totals.get('little_giant_count', 0)} 家"),
            ("产业链完整度", f"{metrics.get('completeness_score', 0)} 分"),
            ("本地配套率", f"{metrics.get('local_support_rate', 0)}%"),
        ]
    )

    # 产业分布表格
    sub_rows = "\n".join(
        f"        <tr><td>{k}</td><td>{v} 家</td></tr>"
        for k, v in metrics.get("sub_industry_distribution", {}).items()
    )
    chain_rows = "\n".join(
        f"        <tr><td>{k}</td><td>{v} 家</td></tr>"
        for k, v in metrics.get("chain_distribution", {}).items()
    )

    # 梯队
    tier_order = ["链主企业", "骨干企业", "高新技术企业", "科技型中小企业", "配套服务企业"]
    tier_rows = "\n".join(
        f"        <tr><td>{k}</td><td>{metrics.get('tier_distribution', {}).get(k, 0)} 家</td></tr>"
        for k in tier_order
        if metrics.get("tier_distribution", {}).get(k, 0) > 0
    )

    # 强弱分析
    def _status_card(title: str, key: str) -> str:
        items = metrics.get("segment_strength", {}).get(key, [])
        content = "、".join(items) if items else "暂无"
        return f"""
    <div class="status-card status-{key}">
      <div class="status-title">{title}</div>
      <p>{content}</p>
    </div>"""

    status_cards = "\n".join([
        _status_card("强势环节", "strong"),
        _status_card("薄弱环节", "weak"),
        _status_card("缺失环节", "missing"),
        _status_card("风险环节", "risk"),
    ])

    segment_rows = "\n".join(
        f"        <tr><td>{row['环节']}</td><td>{row['层级']}</td><td>{row['企业数']}</td><td>{row['年产值']}</td><td><span class='badge badge-{row['状态']}'>{row['状态']}</span></td></tr>"
        for row in _segment_status_table(metrics)
    )

    # 建议
    def _list_items(items: List[str]) -> str:
        if not items:
            return "<p>暂无</p>"
        return "\n      ".join(f"<li>{item}</li>" for item in items)

    top_rows = "\n".join(
        f"        <tr><td>{i}</td><td>{ent.get('name', '—')}</td><td>{ent.get('sub_industry', '—')}</td><td>{ent.get('annual_revenue', 0):.2f}</td></tr>"
        for i, ent in enumerate(metrics.get("top_enterprises", [])[:10], 1)
    )

    html = f"""
<!DOCTYPE html>
<html lang="zh-CN">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{park_name} 产业诊断报告</title>
  <style>
    :root {{
      --bg: #f5f5f7;
      --card: #ffffff;
      --text: #1d1d1f;
      --muted: #6e6e73;
      --border: rgba(0,0,0,0.08);
      --blue: #0056b3;
      --blue-light: rgba(0,86,179,0.12);
      --green: #34c759;
      --green-light: rgba(52,199,89,0.12);
      --orange: #ff9500;
      --orange-light: rgba(255,149,0,0.12);
      --red: #ff3b30;
      --red-light: rgba(255,59,48,0.12);
    }}
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "SF Pro Text", "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
      background: var(--bg);
      color: var(--text);
      line-height: 1.6;
      margin: 0;
      padding: 2rem 1rem;
    }}
    .container {{
      max-width: 960px;
      margin: 0 auto;
    }}
    .header {{
      text-align: center;
      margin-bottom: 2rem;
    }}
    .header h1 {{
      font-size: 2.2rem;
      margin: 0 0 0.5rem;
      letter-spacing: -0.02em;
    }}
    .header .meta {{
      color: var(--muted);
      font-size: 0.95rem;
    }}
    .card {{
      background: var(--card);
      border: 1px solid var(--border);
      border-radius: 20px;
      padding: 1.5rem;
      margin-bottom: 1.25rem;
      box-shadow: 0 2px 8px rgba(0,0,0,0.04);
    }}
    .card h2 {{
      font-size: 1.25rem;
      margin: 0 0 1rem;
      padding-bottom: 0.65rem;
      border-bottom: 1px solid var(--border);
    }}
    .metrics {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(140px, 1fr));
      gap: 1rem;
      margin-bottom: 1.5rem;
    }}
    .metric-card {{
      background: var(--card);
      border: 1px solid var(--border);
      border-radius: 16px;
      padding: 1.1rem;
      text-align: center;
      box-shadow: 0 2px 8px rgba(0,0,0,0.04);
    }}
    .metric-value {{
      font-size: 1.6rem;
      font-weight: 700;
      color: var(--blue);
    }}
    .metric-label {{
      font-size: 0.85rem;
      color: var(--muted);
      margin-top: 0.25rem;
    }}
    .grid-2 {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(320px, 1fr));
      gap: 1rem;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      margin-top: 0.5rem;
    }}
    th, td {{
      padding: 0.6rem 0.75rem;
      border-bottom: 1px solid var(--border);
      text-align: left;
    }}
    th {{
      font-weight: 600;
      background: rgba(0,0,0,0.02);
    }}
    .status-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
      gap: 1rem;
      margin-bottom: 1rem;
    }}
    .status-card {{
      border-radius: 12px;
      padding: 1.1rem 1.25rem;
      border: 1px solid var(--border);
      border-left: 5px solid var(--blue);
      background: var(--card);
    }}
    .status-strong {{ border-left-color: var(--blue); background: var(--blue-light); }}
    .status-weak {{ border-left-color: var(--orange); background: var(--orange-light); }}
    .status-missing {{ border-left-color: var(--red); background: var(--red-light); }}
    .status-risk {{ border-left-color: var(--red); background: var(--red-light); }}
    .status-title {{ font-weight: 700; margin-bottom: 0.35rem; }}
    .badge {{
      display: inline-block;
      padding: 0.15rem 0.55rem;
      border-radius: 9999px;
      font-size: 0.78rem;
      font-weight: 600;
      color: #fff;
    }}
    .badge-强势 {{ background: var(--blue); }}
    .badge-正常 {{ background: var(--green); }}
    .badge-薄弱 {{ background: var(--orange); }}
    .badge-缺失 {{ background: var(--red); }}
    .insight-list {{ margin: 0; padding-left: 1.2rem; }}
    .insight-list li {{ margin-bottom: 0.4rem; }}
    .risk-card {{
      background: var(--red-light);
      border: 1px solid rgba(255, 59, 48, 0.3);
      border-radius: 12px;
      padding: 1rem 1.25rem;
    }}
    .chart {{
      margin: 1rem 0;
      text-align: center;
    }}
    .footer {{
      text-align: center;
      color: var(--muted);
      font-size: 0.85rem;
      margin-top: 2rem;
    }}
    @media print {{
      body {{ background: #fff; padding: 0; }}
      .card {{ break-inside: avoid; }}
    }}
  </style>
</head>
<body>
  <div class="container">
    <div class="header">
      <h1>{park_name} 产业诊断报告</h1>
      <div class="meta">报告生成日期：{date_str}</div>
    </div>

    <div class="metrics">
{metric_cards}
    </div>

    <div class="card">
      <h2>一、产业分布与产业链层级</h2>
      <div class="grid-2">
        <div>
          <h3>产业领域分布</h3>
          <div class="chart">{_img('industry_pie', '产业领域分布')}</div>
          <table>
            <thead><tr><th>产业领域</th><th>企业数</th></tr></thead>
            <tbody>
{sub_rows}
            </tbody>
          </table>
        </div>
        <div>
          <h3>产业链层级分布</h3>
          <div class="chart">{_img('chain_layer', '产业链层级分布')}</div>
          <table>
            <thead><tr><th>层级</th><th>企业数</th></tr></thead>
            <tbody>
{chain_rows}
            </tbody>
          </table>
        </div>
      </div>
    </div>

    <div class="card">
      <h2>二、企业梯队金字塔</h2>
      <div class="chart">{_img('tier_pyramid', '企业梯队金字塔')}</div>
      <table>
        <thead><tr><th>梯队</th><th>企业数</th></tr></thead>
        <tbody>
{tier_rows}
        </tbody>
      </table>
    </div>

    <div class="card">
      <h2>三、产业链强弱分析</h2>
      <div class="status-grid">
{status_cards}
      </div>
      <div class="chart">{_img('segment_strength', '产业链环节布局与强度')}</div>
      <table>
        <thead><tr><th>环节</th><th>层级</th><th>企业数</th><th>年产值</th><th>状态</th></tr></thead>
        <tbody>
{segment_rows}
        </tbody>
      </table>
    </div>

    <div class="card">
      <h2>四、本地配套率与产业链完整度</h2>
      <p><strong>本地配套率：</strong>{metrics.get('local_support_rate', 0)}%</p>
      <p><strong>产业链完整度评分：</strong>{metrics.get('completeness_score', 0)} 分</p>
    </div>

    <div class="card">
      <h2>五、发展建议</h2>
      <h3>整体判断</h3>
      <p>{diagnosis.get('overall_assessment', '暂无')}</p>

      <div class="grid-2">
        <div>
          <h3>核心优势</h3>
          <ul class="insight-list">{_list_items(diagnosis.get('core_strengths', []))}</ul>
        </div>
        <div>
          <h3>关键短板</h3>
          <ul class="insight-list">{_list_items(diagnosis.get('key_gaps', []))}</ul>
        </div>
      </div>

      <h3>招商补链建议</h3>
      <ul class="insight-list">{_list_items(diagnosis.get('investment_suggestions', []))}</ul>

      <h3>企业梯度培育建议</h3>
      <ul class="insight-list">{_list_items(diagnosis.get('cultivation_suggestions', []))}</ul>

      <h3>政策建议</h3>
      <ul class="insight-list">{_list_items(diagnosis.get('policy_suggestions', []))}</ul>

      <h3>风险提醒</h3>
      <div class="risk-card">{diagnosis.get('risk_warning', '暂无')}</div>
    </div>

    <div class="card">
      <h2>六、TOP10 头部企业</h2>
      <div class="chart">{_img('top_enterprises', 'TOP10 头部企业')}</div>
      <table>
        <thead><tr><th>排名</th><th>企业名称</th><th>产业领域</th><th>年产值（亿元）</th></tr></thead>
        <tbody>
{top_rows}
        </tbody>
      </table>
    </div>

    <div class="footer">本报告由园区产业分析智能体自动生成</div>
  </div>
</body>
</html>
"""
    return html


# ==================== 内部辅助函数 ====================


def _add_chart_or_placeholder(doc, image_bytes: Optional[bytes], title: Optional[str] = None):
    """向 Word 文档插入图片或占位文字"""
    from docx.shared import Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    if title:
        # 标题已在调用处添加
        pass
    if image_bytes:
        image_stream = io.BytesIO(image_bytes)
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        run.add_picture(image_stream, width=Inches(5.5))
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    else:
        doc.add_paragraph("（未提供图表）")


def _set_docx_cell_shading(cell, hex_color: str):
    """设置 docx 单元格背景色"""
    from docx.oxml import parse_xml
    from docx.shared import RGBColor

    fill = hex_color.lstrip("#")
    shading_elm = parse_xml(
        '<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'w:fill="{fill}"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading_elm)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(255, 255, 255)


def _hex_to_rgb(hex_color: str):
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def _pdf_safe_multi_cell(pdf, text: str, line_height: float = 5.5):
    """重置 x 到左边距后渲染多行文本"""
    pdf.set_xy(pdf.l_margin, pdf.get_y())
    content_width = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.multi_cell(content_width, line_height, str(text))
    pdf.set_x(pdf.l_margin)


def _pdf_table_row(
    pdf,
    cells,
    col_widths,
    line_height: float,
    aligns=None,
    fills=None,
):
    if aligns is None:
        aligns = ["L"] * len(cells)
    if fills is None:
        fills = [None] * len(cells)

    x_start = pdf.get_x()
    y_start = pdf.get_y()

    heights = []
    for text, width in zip(cells, col_widths):
        try:
            lines = pdf.multi_cell(width, line_height, str(text), split_only=True)
            heights.append(len(lines) * line_height)
        except Exception:
            heights.append(line_height)
    row_height = max(heights + [line_height])

    for i, (text, width, align, fill) in enumerate(zip(cells, col_widths, aligns, fills)):
        pdf.set_xy(x_start + sum(col_widths[:i]), y_start)
        if fill:
            pdf.set_fill_color(*fill)
            pdf.set_text_color(255, 255, 255)
            pdf.multi_cell(width, line_height, str(text), border=1, align=align, fill=True)
            pdf.set_text_color(0, 0, 0)
        else:
            pdf.multi_cell(width, line_height, str(text), border=1, align=align)

    pdf.set_xy(x_start, y_start + row_height)
    return row_height


def _pdf_insert_chart(pdf, image_bytes: Optional[bytes], title: Optional[str] = None):
    if title:
        pdf.set_font("cn", "B", 12)
        pdf.cell(0, 8, title, ln=True)
        pdf.ln(1)
    if image_bytes:
        image_stream = io.BytesIO(image_bytes)
        pdf.image(image_stream, x=20, w=170)
        pdf.ln(5)
    else:
        pdf.set_font("cn", "", 10)
        _pdf_safe_multi_cell(pdf, "（未提供图表）")
        pdf.ln(2)
