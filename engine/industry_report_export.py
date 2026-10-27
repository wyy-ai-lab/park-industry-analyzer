"""
产业梳理报告导出：Word / PDF
==============================
与 engine.industry_analyzer.to_markdown 保持同一内容结构：
产业分类 → 环节梳理 → 强链补链延链 → 招商目标清单 → 研究院/协会资源。

中文字体策略与 engine.report_export / engine.park_report_export 一致：
PDF 使用 fpdf2 + 系统/缓存 CJK 字体；Word 使用 python-docx 并设置东亚字体。
"""

import io
from datetime import datetime
from typing import Any, Dict, List, Tuple

from engine.report_export import _find_cjk_font
from engine.park_report_export import _pdf_safe_multi_cell, _pdf_table_row
from engine.industry_analyzer import _report_summary, _report_disclaimer


def _today() -> str:
    return datetime.now().strftime("%Y年%m月%d日")


def _mode_text(result: Dict[str, Any]) -> str:
    return ("LLM 智能梳理" if result.get("mode") == "llm"
            else "模板模式（未配置 API Key，仅基础归位）")


def _source_text(t: Dict[str, Any]) -> str:
    return "知识库·已核实" if t.get("source") == "知识库" else "AI生成·待核实"


def _targets_table(result: Dict[str, Any]) -> Tuple[List[str], List[List[str]]]:
    headers = ["目标企业", "对应环节", "所在地", "配套逻辑", "落地线索", "来源"]
    rows = [[t.get("name", ""), t.get("segment", ""), t.get("region", ""),
             t.get("reason", ""), t.get("signal", ""), _source_text(t)]
            for t in result.get("targets", [])]
    return headers, rows


# ---------------------------------------------------------------- Word

def build_industry_word_report(result: Dict[str, Any]) -> io.BytesIO:
    """生成产业梳理 Word 报告，返回 BytesIO"""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(10.5)

    # 封面
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title.add_run("园区产业梳理报告")
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(29, 29, 31)
    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    info.add_run(f"生成模式：{_mode_text(result)}    报告日期：{_today()}").font.size = Pt(12)
    if result.get("focus_hint") or result.get("region"):
        extra = doc.add_paragraph()
        extra.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        parts = []
        if result.get("focus_hint"):
            parts.append(f"产业方向提示：{result['focus_hint']}")
        if result.get("region"):
            parts.append(f"目标区域偏好：{result['region']}")
        extra.add_run("　".join(parts)).font.size = Pt(12)
    doc.add_page_break()

    # 摘要
    doc.add_heading("摘  要", level=1)
    for line in _report_summary(result):
        # 去掉 Markdown 加粗记号
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(line.replace("**", ""))
    doc.add_page_break()

    # 一、产业分类
    doc.add_heading("一、产业分类", level=1)
    for group in result["classification"].get("industries", []):
        doc.add_heading(group["industry"], level=2)
        for e in group.get("enterprises", []):
            desc = "、".join(x for x in (e.get("products"), e.get("tech"))
                            if x and x != "未提及")
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(e.get("name", ""))
            r.font.bold = True
            if desc:
                p.add_run(f"：{desc}")

    # 二、上中下游环节梳理
    doc.add_heading("二、产业链上下游分析", level=1)
    for layer in ("上游", "中游", "下游"):
        doc.add_heading(layer, level=2)
        segs = [s for s in result["chain_map"].get("segments", [])
                if s.get("layer") == layer]
        if not segs:
            doc.add_paragraph("（无企业布局）")
        for s in segs:
            names = "、".join(s.get("enterprises", [])) or "（缺失）"
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(s.get("segment", "")).font.bold = True
            p.add_run(f"：{names}")

    # 三、强链补链延链路径
    doc.add_heading("三、产业发展建议", level=1)
    paths = result.get("paths", {})
    for title, key in (("强链", "strengthen"), ("补链", "complete"), ("延链", "extend")):
        doc.add_heading(title, level=2)
        items = paths.get(key, [])
        if not items:
            doc.add_paragraph("（无）")
        for item in items:
            seg = item.get("segment") or item.get("direction", "")
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(seg).font.bold = True
            p.add_run(f"：{item.get('action', '')}（{item.get('rationale', '')}）")

    # 四、招商目标企业清单
    doc.add_heading("四、招商目标企业清单", level=1)
    if result.get("targets"):
        headers, rows = _targets_table(result)
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            cell.paragraphs[0].runs[0].font.bold = True
        for row in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = text
        doc.add_paragraph()
        note = doc.add_paragraph()
        note.add_run(
            "注：知识库条目为人工整理、已核实，对接前建议复核最新动态；"
            "AI 生成条目请在正式对接前逐一核实企业真实性与落地意愿。").font.size = Pt(9)
    else:
        doc.add_paragraph("（未命中知识库且未配置 API Key，无目标企业清单）")

    # 五、研究院 / 协会资源
    doc.add_heading("五、研究院 / 协会资源", level=1)
    res = result.get("resources", {})
    if res:
        doc.add_heading("研究机构", level=2)
        for i in res.get("institutes", []):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(i.get("name", "")).font.bold = True
            p.add_run(f"（{i.get('location', '')}）：{i.get('focus', '')}")
        doc.add_heading("行业组织", level=2)
        for a in res.get("associations", []):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(a.get("name", "")).font.bold = True
            p.add_run(f"：{a.get('focus', '')}")
    else:
        doc.add_paragraph("（未生成产业资源清单）")

    # 六、使用说明
    doc.add_heading("六、使用说明", level=1)
    doc.add_paragraph(_report_disclaimer(result))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- PDF

def build_industry_pdf_report(result: Dict[str, Any]) -> io.BytesIO:
    """生成产业梳理 PDF 报告，返回 BytesIO"""
    from fpdf import FPDF

    font_path = _find_cjk_font()
    if not font_path:
        raise RuntimeError(
            "未找到中文字体，PDF 导出无法显示中文。"
            "请在 Windows 系统字体目录保留 simsun、simhei 或 msyh 字体之一。"
        )

    pdf = FPDF()
    pdf.add_font("cn", "", font_path, uni=True)
    pdf.add_font("cn", "B", font_path, uni=True)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_left_margin(15)
    pdf.set_right_margin(15)

    # 封面
    pdf.add_page()
    pdf.set_font("cn", "B", 26)
    pdf.cell(0, 24, "园区产业梳理报告", ln=True, align="C")
    pdf.ln(6)
    pdf.set_font("cn", "", 13)
    pdf.cell(0, 9, f"生成模式：{_mode_text(result)}", ln=True, align="C")
    pdf.cell(0, 9, f"报告日期：{_today()}", ln=True, align="C")
    if result.get("focus_hint"):
        pdf.cell(0, 9, f"产业方向提示：{result['focus_hint']}", ln=True, align="C")
    if result.get("region"):
        pdf.cell(0, 9, f"目标区域偏好：{result['region']}", ln=True, align="C")

    # 摘要
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "摘  要", ln=True)
    pdf.ln(1)
    pdf.set_font("cn", "", 10)
    for line in _report_summary(result):
        _pdf_safe_multi_cell(pdf, "· " + line.replace("**", ""), line_height=5.8)
    pdf.ln(2)

    # 一、产业分类
    pdf.add_page()
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "一、产业分类", ln=True)
    pdf.ln(1)
    for group in result["classification"].get("industries", []):
        pdf.set_font("cn", "B", 12)
        pdf.cell(0, 8, str(group["industry"]), ln=True)
        pdf.set_font("cn", "", 10)
        for e in group.get("enterprises", []):
            desc = "、".join(x for x in (e.get("products"), e.get("tech"))
                            if x and x != "未提及")
            _pdf_safe_multi_cell(pdf, f"· {e.get('name', '')}" + (f"：{desc}" if desc else ""),
                                 line_height=5.8)
        pdf.ln(2)

    # 二、上中下游环节梳理
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "二、产业链上下游分析", ln=True)
    pdf.ln(1)
    for layer in ("上游", "中游", "下游"):
        pdf.set_font("cn", "B", 12)
        pdf.cell(0, 8, layer, ln=True)
        pdf.set_font("cn", "", 10)
        segs = [s for s in result["chain_map"].get("segments", [])
                if s.get("layer") == layer]
        if not segs:
            _pdf_safe_multi_cell(pdf, "（无企业布局）", line_height=5.8)
        for s in segs:
            names = "、".join(s.get("enterprises", [])) or "（缺失）"
            _pdf_safe_multi_cell(pdf, f"· {s.get('segment', '')}：{names}", line_height=5.8)
        pdf.ln(2)

    # 三、强链补链延链路径
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "三、产业发展建议", ln=True)
    pdf.ln(1)
    paths = result.get("paths", {})
    for title, key in (("强链", "strengthen"), ("补链", "complete"), ("延链", "extend")):
        pdf.set_font("cn", "B", 12)
        pdf.cell(0, 8, title, ln=True)
        pdf.set_font("cn", "", 10)
        items = paths.get(key, [])
        if not items:
            _pdf_safe_multi_cell(pdf, "（无）", line_height=5.8)
        for item in items:
            seg = item.get("segment") or item.get("direction", "")
            _pdf_safe_multi_cell(
                pdf, f"· {seg}：{item.get('action', '')}（{item.get('rationale', '')}）",
                line_height=5.8)
        pdf.ln(2)

    # 四、招商目标企业清单
    if result.get("targets"):
        pdf.add_page()
        pdf.set_font("cn", "B", 16)
        pdf.cell(0, 10, "四、招商目标企业清单", ln=True)
        pdf.ln(1)
        headers, rows = _targets_table(result)
        widths = [34, 20, 24, 40, 34, 24]
        pdf.set_font("cn", "B", 8.5)
        _pdf_table_row(pdf, headers, widths, 5.5,
                       aligns=["C"] * len(headers), fills=[(0, 82, 217)] * len(headers))
        pdf.set_font("cn", "", 8.5)
        for row in rows:
            _pdf_table_row(pdf, row, widths, 5.5, aligns=["L"] * len(row))
        pdf.ln(3)
        pdf.set_font("cn", "", 9)
        _pdf_safe_multi_cell(
            pdf,
            "注：知识库条目为人工整理、已核实，对接前建议复核最新动态；"
            "AI 生成条目请在正式对接前逐一核实企业真实性与落地意愿。",
            line_height=5)

    # 五、研究院 / 协会资源
    pdf.ln(2)
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "五、研究院 / 协会资源", ln=True)
    pdf.ln(1)
    res = result.get("resources", {})
    if res:
        pdf.set_font("cn", "B", 12)
        pdf.cell(0, 8, "研究机构", ln=True)
        pdf.set_font("cn", "", 10)
        for i in res.get("institutes", []):
            _pdf_safe_multi_cell(
                pdf, f"· {i.get('name', '')}（{i.get('location', '')}）：{i.get('focus', '')}",
                line_height=5.8)
        pdf.ln(2)
        pdf.set_font("cn", "B", 12)
        pdf.cell(0, 8, "行业组织", ln=True)
        pdf.set_font("cn", "", 10)
        for a in res.get("associations", []):
            _pdf_safe_multi_cell(pdf, f"· {a.get('name', '')}：{a.get('focus', '')}",
                                 line_height=5.8)
    else:
        pdf.set_font("cn", "", 10)
        _pdf_safe_multi_cell(pdf, "（未生成产业资源清单）", line_height=5.8)

    # 六、使用说明
    pdf.ln(2)
    pdf.set_font("cn", "B", 16)
    pdf.cell(0, 10, "六、使用说明", ln=True)
    pdf.ln(1)
    pdf.set_font("cn", "", 10)
    _pdf_safe_multi_cell(pdf, _report_disclaimer(result), line_height=5.8)

    return io.BytesIO(pdf.output())
